import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import win32com.client
import pymupdf

def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    for s in tcPr.xpath('w:shd'):
        tcPr.remove(s)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_vcenter(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    for v in tcPr.xpath('w:vAlign'):
        tcPr.remove(v)
    vAlign = OxmlElement('w:vAlign')
    vAlign.set(qn('w:val'), 'center')
    tcPr.append(vAlign)

def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    cantSplit = OxmlElement('w:cantSplit')
    trPr.append(cantSplit)

def set_callout_box(cell, bg_color='F2F4F8', border_color='003366'):
    tcPr = cell._tc.get_or_add_tcPr()
    for s in tcPr.xpath('w:shd'):
        tcPr.remove(s)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
    tcPr.append(shd)
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'  <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    set_cell_margins(cell, top=80, bottom=80, left=120, right=100)

def set_code_box(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    for s in tcPr.xpath('w:shd'):
        tcPr.remove(s)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8F9FA"/>')
    tcPr.append(shd)
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>\n'
        f'  <w:left w:val="single" w:sz="18" w:space="0" w:color="0969DA"/>\n'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>\n'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    set_cell_margins(cell, top=60, bottom=60, left=90, right=90)

def style_run(run, font_name='Arial', size_pt=10.0, bold=False, italic=False, color_rgb=(0x3B, 0x38, 0x38)):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = RGBColor(*color_rgb)

def convert_docx_to_pdf(docx_path, pdf_path):
    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    try:
        doc = word.Documents.Open(os.path.abspath(docx_path))
        doc.SaveAs(os.path.abspath(pdf_path), FileFormat=17) # wdFormatPDF = 17
        doc.Close()
    finally:
        word.Quit()
    print(f"Exported Master PDF successfully: {pdf_path}")

def build_master_guide():
    doc = docx.Document()

    # Set Margins (A4 standard)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.8)

    def add_p(text='', bold=False, italic=False, size_pt=10.0, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=3, left_indent=0, color_rgb=(0x3B, 0x38, 0x38), keep_with_next=False):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.keep_with_next = keep_with_next
        if left_indent > 0:
            p.paragraph_format.left_indent = Inches(left_indent)
        if text:
            r = p.add_run(text)
            style_run(r, font_name='Arial', size_pt=size_pt, bold=bold, italic=italic, color_rgb=color_rgb)
        return p

    def add_h1(text):
        return add_p(text, bold=True, size_pt=14.0, space_before=14, space_after=4, color_rgb=(0, 51, 102), keep_with_next=True)

    def add_h2(text):
        return add_p(text, bold=True, size_pt=11.5, space_before=10, space_after=3, color_rgb=(0, 0, 0), keep_with_next=True)

    def add_h3(text):
        return add_p(text, bold=True, size_pt=10.5, space_before=6, space_after=2, color_rgb=(0x20, 0x20, 0x20), keep_with_next=True)

    def add_callout(text, bold_prefix='', bg_color='F2F4F8', border_color='003366'):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.6)
        set_callout_box(cell, bg_color=bg_color, border_color=border_color)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.12
        if bold_prefix:
            r1 = p.add_run(bold_prefix + ' ')
            style_run(r1, font_name='Arial', size_pt=9.5, bold=True, color_rgb=(0, 51, 102))
        r2 = p.add_run(text)
        style_run(r2, font_name='Arial', size_pt=9.5, bold=False, color_rgb=(0x22, 0x22, 0x22))
        add_p(space_after=2)

    def add_code(text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.6)
        set_code_box(cell)
        lines = text.strip().split('\n')
        cell.paragraphs[0].text = ''
        for i, line in enumerate(lines):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(line)
            style_run(r, font_name='Consolas', size_pt=8.5, bold=False, color_rgb=(0x1F, 0x23, 0x28))
        add_p(space_after=2)

    # --- COVER / TITLE BLOCK ---
    add_p('PT NIRWANA BHUMI ENERGY', bold=True, size_pt=18.0, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=10, space_after=4, color_rgb=(0, 0, 0))
    add_p('DOCUMENT CONTROL CENTER (DCC) & SYSTEM ENGINEERING', bold=True, size_pt=11.0, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6, color_rgb=(0, 51, 102))
    add_p('PANDUAN LENGKAP ARSITEKTUR FORMATTING, MASTER PROMPT SISTEM WEB SURAT, & SPESIFIKASI ENGINE MULTIBAHASA', bold=True, size_pt=15.0, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4, color_rgb=(0, 0, 0))
    add_p('Implementasi Terintegrasi pada Antigravity IDE untuk Otomasi Editing Dokumen, Analisis Folder, Pilihan Bahasa (ID / EN / Bilingual), dan Sinkronisasi Cloud Google Drive', bold=False, italic=True, size_pt=10.0, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12, color_rgb=(0x55, 0x55, 0x55))

    # Meta Table
    t_meta = doc.add_table(rows=4, cols=2)
    t_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ('Klasifikasi Dokumen', 'Spesifikasi Sistem Teknis & Master AI Agent Prompt'),
        ('Target Platform', 'Antigravity IDE (Web Project "Surat"), Google Docs MCP, Python-Docx Engine'),
        ('Penyusun & Otorisasi', 'Mazmur Gusti Agung Larosa (Lead Document Controller & Finance/Eng Director)'),
        ('Tanggal Terbit / Status', '17 Oktober 2026 / Versi 1.0 (Final & Siap Eksekusi)')
    ]
    for r_i, (k, v) in enumerate(meta_data):
        row = t_meta.rows[r_i]
        set_row_cant_split(row)
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = Inches(2.2), Inches(4.4)
        c0.paragraphs[0].text = k
        c1.paragraphs[0].text = v
        set_cell_shading(c0, 'F2F4F8')
        set_cell_shading(c1, 'FFFFFF')
        set_cell_vcenter(c0)
        set_cell_vcenter(c1)
        set_cell_margins(c0, top=40, bottom=40, left=60, right=60)
        set_cell_margins(c1, top=40, bottom=40, left=60, right=60)
        style_run(c0.paragraphs[0].runs[0], font_name='Arial', size_pt=9.0, bold=True, color_rgb=(0, 0, 0))
        style_run(c1.paragraphs[0].runs[0], font_name='Arial', size_pt=9.0, bold=False, color_rgb=(0x3B, 0x38, 0x38))

    add_p(space_after=8)

    # --- BAGIAN 1 ---
    add_h1('BAGIAN 1: RINGKASAN EKSEKUTIF & ARSITEKTUR SISTEM')
    add_p('Dokumen ini disusun sebagai blueprint komprehensif bagi sistem web otomasi persuratan dan kontrol dokumen yang sedang dikembangkan di Antigravity IDE (proyek "surat"). Sistem ini bertujuan mengotomatiskan penerimaan, konversi, validasi, pengeditan tata letak baku, penomoran kode unik 6-segmen, serta pengarsipan cloud seluruh dokumen resmi PT Nirwana Bhumi Energy (NBE).')
    
    add_callout(
        'Prinsip Utama: Sistem web yang dibangun di Antigravity IDE wajib mampu mereplikasi 100% struktur visual, tata letak, hierarki halaman, dan aturan sub-pixel yang telah berhasil dibuktikan pada berkas acuan (0002-KB-GNR-CEO-X-2026). Tidak boleh ada deviasi format, tumpahan halaman liar, atau pemotongan tabel canggung.',
        bold_prefix='PANDUAN MUTU UTAMA:'
    )

    add_h2('1.1 Alur Kerja Sistem Web Terintegrasi (End-to-End Pipeline)')
    add_p('Sistem bekerja melalui 4 tahapan eksekusi logis yang saling terhubung:')
    add_p('1.  Tahap Ingesti & Pemisahan Berkas: Memindai folder input, memisahkan berkas Word (.docx) dan PDF (.pdf), serta mengeksekusi konversi otomatis PDF-ke-Word jika isi berkas berbeda.', left_indent=0.3)
    add_p('2.  Tahap Analisis & Dashboard: Menghitung statistik berkas (Total dokumen, jumlah PDF, jumlah Word, dan jumlah hasil konversi baru) serta menampilkan ringkasan kesiapan berkas.', left_indent=0.3)
    add_p('3.  Tahap Standarisasi Template & Editor Interaktif: Menggunakan file acuan (.txt untuk struktur teks dan .pdf untuk visual looks), menghubungkan Google Docs MCP / Python DOCX engine, menampilkan Preview nyata (real paper-size preview) tepat di bawah tab penomoran, serta menyediakan form klasifikasi dan TOMBOL PILIHAN BAHASA.', left_indent=0.3)
    add_p('4.  Tahap Routing Otomatis & Pengarsipan Google Drive: Menghasilkan kode unik 6 segmen, mengunci kode via Tombol "Fix Kode", membuat pohon direktori berjenjang otomatis berdasarkan klasifikasi surat, dan menyinkronkan berkas ke folder induk Google Drive ID: 1ny_1VhXfSaNrj_XdoOsC7V7BO-z0K8eF.', left_indent=0.3)

    # --- BAGIAN 2 ---
    add_h1('BAGIAN 2: MASTER PROMPT TERINTEGRASI UNTUK ANTIGRAVITY IDE')
    add_p('Berikut adalah naskah Master Prompt yang telah disempurnakan dari file "prom surat.docx". Prompt ini dirancang dengan standar Agentic Instruction bertaraf internasional, memiliki pembagian fase yang ketat, instruksi error-handling, serta siap disalin (copy-paste) ke dalam sistem Antigravity IDE Anda:')

    master_prompt_text = (
        '=== MASTER SYSTEM PROMPT: AUTONOMOUS DOCUMENT MANAGEMENT & GENERATION AGENT ===\n'
        'Anda adalah AI Agent Document Engineering & Web System Architect untuk proyek "Surat" PT Nirwana Bhumi Energy (NBE) di Antigravity IDE.\n'
        'Tugas Anda adalah memproses unggahan dokumen korporat, melakukan verifikasi dan standardisasi isi, menyediakan interface editor berbasis Google Docs MCP / python-docx, mengarahkan tata letak halaman sesuai standar IMS NBE, serta mengarsipkan dokumen ke Google Drive secara otomatis.\n\n'
        'JALANKAN SISTEM MELALUI 4 TAHAPAN KERJA BERIKUT SECARA DISIPLIN:\n\n'
        '--- TAHAP 1: INGESTI FOLDER, PEMISAHAN BERKAS, & AUTO-CONVERT PDF ---\n'
        '1. Sistem menerima input unggahan berupa satu folder yang berisi kumpulan berkas (contoh: 10 berkas campuran Word dan PDF).\n'
        '2. Ketika pengguna menekan tombol "Next":\n'
        '   a. Sistem memindai seluruh berkas dan otomatis membuat 2 subfolder kerja: "/word" dan "/pdf".\n'
        '   b. Jika sebuah dokumen memiliki versi Word dan PDF dengan isi yang identik, pisahkan masing-masing berkas ke folder yang sesuai.\n'
        '   c. Jika ditemukan berkas PDF yang isinya berbeda atau belum memiliki pasangan file Word, sistem secara otomatis mengeksekusi konversi PDF ke DOCX (menggunakan PyMuPDF / pdf2docx / OCR parser), lalu memasukkan file Word baru tersebut ke dalam folder "/word".\n'
        '   d. Prioritas proses selanjutnya 100% diarahkan pada berkas-berkas di folder "/word", sementara berkas PDF diarsipkan sebagai referensi audit.\n\n'
        '--- TAHAP 2: STATISTICAL REPORTING & DASHBOARD ANALISIS ---\n'
        '1. Tampilkan kartu ringkasan analitik dokumen secara real-time pada dashboard:\n'
        '   - Jumlah Total Dokumen yang Diterima Sistem\n'
        '   - Jumlah Dokumen Asli Format PDF\n'
        '   - Jumlah Dokumen Asli Format WORD\n'
        '   - Jumlah Dokumen Word Baru Hasil Konversi (Tampilkan angka riil; jika tidak ada, tulis "0")\n'
        '2. Tampilkan tombol "Next" untuk mengarahkan pengguna ke tahap standarisasi template.\n\n'
        '--- TAHAP 3: ACUAN TEMPLATE DOKUMEN, PREVIEW REAL-SIZE, & ENGINE EDITOR ---\n'
        '1. Buat Tab Baru bernama "Acuan Template Dokumen" di bawah tab "Kelola Master Kategori & Kode Surat". Tab ini memuat:\n'
        '   a. Kolom Unggah Berkas Acuan (.txt): Berfungsi sebagai dasar aturan teks, diktum, dan hierarki logika utama dokumen.\n'
        '   b. Kolom Unggah Berkas Acuan (.pdf): Berfungsi sebagai dasar geometri visual, posisi logo, margin, kop, dan estetika layout.\n'
        '2. Integrasi Engine Editor Dokumen (Google Docs MCP & Python Engine):\n'
        '   - Sistem mengonversi file Word menjadi dokumen terkontrol yang dapat diedit secara programmatic via Google Docs MCP / python-docx.\n'
        '   - Tampilkan daftar berkas Word yang diterima lengkap dengan tombol "Edit" di setiap baris.\n'
        '3. Komponen Real-Time Document Preview (Real Paper Size):\n'
        '   - LETAKKAN PREVIEW TEPAT DI BAWAH tab "Penomoran Kode Dokumen & Klasifikasi NBE" (DILARANG menempatkannya di samping karena mengganggu ergonomi pandangan).\n'
        '   - Preview wajib menampilkan render nyata skala kertas A4 (Halaman 1 Cover, Halaman 2 Revision Index, Halaman 3-4 Isi Surat).\n'
        '4. Form Kontrol & Parameterisasi Surat:\n'
        '   - Field Ubah Judul Dokumen (Otomatis menyinkronkan Judul pada Cover dan Judul pada Kop Halaman 2 ke atas).\n'
        '   - Form Tabel Revision Index:\n'
        '     * Kolom Rev (dimulai dari 0 untuk edisi awal).\n'
        '     * Kolom Date (tanggal penetapan dokumen).\n'
        '     * Kolom Description (Otomatis: "Initial Edition – [Judul Dokumen]" untuk revisi 0, atau deskripsi rincian perubahan untuk revisi lanjutan).\n'
        '     * Kolom Prepared By (Nama penyusun dokumen).\n'
        '     * CATATAN: Bagian "Catatan Kendali Dokumen" DITIADAKAN agar Halaman 2 bersih.\n'
        '   - Dropdown Klasifikasi Kode Surat 6-Segmen:\n'
        '     * Dropdown 1: Pilih Kategori Surat (Document Type: KB, SOP, IK, FORM, CSMS, dll.)\n'
        '     * Dropdown 2: Pilih Sub-Kategori Surat (Kode ZZ, misal GP = General Procedure)\n'
        '     * Dropdown 3: Pilih Sub-sub Kategori Surat (Kode AA, misal AI = Audit Internal)\n'
        '     * Dropdown 4: Pilih Penerbit Surat (Division, misal GNR, CEO, ENG, OPR)\n'
        '     * Dropdown 5: Pilih Pejabat Penyetuju (Approver: CEO / Bima Satryo Hartanto)\n'
        '5. TOMBOL PILIHAN BAHASA DOKUMEN (LANGUAGE SELECTOR):\n'
        '   - Sediakan tombol pemilihan bahasa dokumen dengan 3 opsi:\n'
        '     [ Bahasa Indonesia ]  |  [ British English ]  |  [ Bilingual (ID - EN-GB) ]\n'
        '   - Ketika opsi dipilih, engine otomatis menyusun cover, kop, tabel revisi, dan isi surat sesuai standar bahasa tersebut secara konsisten 4 halaman.\n'
        '6. Tombol "Simpan & Upload": Memfinalisasi pengeditan seluruh berkas.\n\n'
        '--- TAHAP 4: ROUTING FOLDER OTOMATIS & INTEGRASI GOOGLE DRIVE ---\n'
        '1. Sistem memverifikasi ketersediaan folder induk berdasarkan Kategori Surat (misal: folder "CSMS"). Jika belum ada, sistem menampilkan peringatan panduan bagi pengguna.\n'
        '2. Di dalam folder Kategori, bot otomatis membuat folder baru sesuai nama kepanjangan Sub-Kategori (Kode ZZ, contoh: "General Procedure"). Jika sudah ada, gunakan yang ada.\n'
        '3. Di dalam folder Sub-Kategori, bot membuat folder baru sesuai kepanjangan Sub-sub Kategori (Kode AA, contoh: "Audit Internal").\n'
        '4. Di dalam folder Sub-sub Kategori, bot membuat folder khusus dengan nama Judul Dokumen Lengkap (contoh: "Kebijakan Audit Internal").\n'
        '5. Tombol "Fix Kode": Mengunci nomor unik [NomorUrut4Digit]-[Kategori]-[Divisi]-[Pengesah]-[BulanRomawi]-[Tahun], mencatatkan di Master Register, dan mengunggah dokumen digital final ke Google Drive Korporat (Folder Root ID: 1ny_1VhXfSaNrj_XdoOsC7V7BO-z0K8eF).'
    )
    add_code(master_prompt_text)

    # --- BAGIAN 3 ---
    add_h1('BAGIAN 3: PANDUAN FORMATTING & STANDAR BAKU DOKUMEN 4 HALAMAN')
    add_p('Agar dokumen yang dikeluarkan oleh sistem web Anda sama persis dengan berkas acuan yang baru saja kita selesaikan, engine dokumen wajib mematuhi parameter teknis berikut:')

    add_h2('3.1 Geometri Halaman & Tipografi Dasar')
    add_p('• Ukuran Kertas: A4 Standard (Lebar: 210 mm / 8.27", Tinggi: 297 mm / 11.69").')
    add_p('• Margin Halaman: Kiri = 2.2 cm (0.866"), Kanan = 2.0 cm (0.787"), Atas = 2.0 cm (0.787"), Bawah = 2.0 cm (0.787").')
    add_p('• Keluarga Huruf (Font): Arial untuk seluruh teks dokumen.')
    add_p('• Palet Warna Hex: Hitam Judul (#000000), Abu Gelap Body (#3B3838), Biru Nomor Dokumen (#003366), Header Tabel Biru Muda (#D9E1F2), Alternating Row (#F9FBFD / #FFFFFF), Border Tabel (#BFBFBF).')

    add_h2('3.2 Spesifikasi Halaman 1 (Cover Page)')
    add_p('1.  Logo Perusahaan: Berada tepat di tengah (Centered) dengan lebar 1.85 inci (4.7 cm).')
    add_p('2.  Nama Perusahaan: PT NIRWANA BHUMI ENERGY (Arial 20.0 pt Bold, Centered, #000000).')
    add_p('3.  Judul Dokumen: SURAT KEPUTUSAN DIREKSI (Arial 15-16 pt Bold), Nama Ketetapan (Arial 11-11.5 pt Bold), Nomor Dokumen (Arial 10.5-11 pt Bold, #003366).')
    add_p('4.  Jeda Spasi Vertikal Pendorong: 5 paragraf kosong line spacing 11-13 pt, diikuti 5 paragraf kosong line spacing 2 pt untuk mendorong blok otorisasi ke kuadran bawah tanpa memicu overflow.')
    add_p('5.  BLOK OTORISASI DI UJUNG KANAN BAWAH: Menggunakan left indent 3.4 - 3.65 inci, berisi Approved by (Bima Satryo Hartanto / CEO), Verified by (Dennis Raditya / CTO), dan Prepared by (Mazmur Gusti Agung Larosa / DCC Lead).')
    add_p('6.  Penutup Halaman 1: Diakhiri dengan Explicit Page Break (w:br type="page").')

    add_h2('3.3 Spesifikasi Halaman 2 (Revision Index)')
    add_p('1.  Header Kop Resmi NBE: Terdiri atas 4 baris sel (Division, Category, Title, Rev/Code/Date) dengan logo NBE di sel kiri.')
    add_p('2.  Judul REVISION INDEX Ditebalkan: Wajib menggunakan font Arial 12.0 pt Bold, Centered, spasi sebelum 18 pt, spasi sesudah 12 pt.')
    add_p('3.  Tabel Revision Index: 4 Kolom (Rev, Date, Description, Prepared By). Baris header diberi warna latar #D9E1F2, Arial 9.0-9.5 pt Bold, cell margin top/bottom 70 dxa, left/right 80 dxa, vertical align center.')
    add_p('4.  Catatan Kendali Dokumen: Wajib DIHAPUS SEPENUHNYA sehingga Halaman 2 hanya memuat Tabel Revision Index secara bersih.')
    add_p('5.  Penutup Halaman 2: Diakhiri dengan Explicit Page Break ke Halaman 3.')

    add_h2('3.4 Spesifikasi Halaman 3 (Permulaan Isi Surat / Letter Body)')
    add_p('1.  Pemisahan Halaman: Isi surat resmi dimulai tepat pada HALAMAN BARU (Halaman 3).')
    add_p('2.  Judul Resmi & Konsiderans: Memuat Judul Surat Keputusan, Nomor, Tentang, Opening (DIREKSI PT NIRWANA BHUMI ENERGY), Konsiderans Menimbang (butir a, b, c), dan Mengingat (butir 1 s.d. 5).')
    add_p('3.  Diktum Pembuka & KESATU: MEMUTUSKAN, Menetapkan, serta Diktum KESATU yang mengatur Standar Penomoran 6 Segmen dan Asas Unik Mutlak (1 Kode = 1 Dokumen Sah).')
    add_p('4.  Penutup Halaman 3: Diakhiri dengan Explicit Page Break ke Halaman 4.')

    add_h2('3.5 Spesifikasi Halaman 4 (Lanjutan Diktum, Tim DCC, & Tanda Tangan)')
    add_p('1.  Diktum KEDUA: Memuat penunjukan Tim Document Control Center (DCC).')
    add_p('2.  Tabel Tim DCC (4 Kolom): No, Nama Personel, Jabatan Struktural, Peran & Tanggung Jawab DCC. Menggunakan lebar penuh (6.44 inci), header #D9E1F2, sel data zebra alternating fill (#FFFFFF & #F9FBFD), serta cantSplit pada seluruh baris.')
    add_p('3.  Diktum KETIGA, KEEMPAT, KELIMA: Mengatur tombol Fix Kode & Google Drive korporat, alur revisi terkontrol (-rev1, obsolete), dan klausul masa berlaku.')
    add_p('4.  Klausul Tempat & Tanggal Penetapan: Ditetapkan di Semarang / Tuban, 17 Oktober 2026.')
    add_p('5.  Tabel Tanda Tangan 3 Kolom: Terdiri dari 3 baris: Baris Jabatan Pengesah, Baris Ruang Tanda Tangan/Meterai ([ Meterai Rp 10.000 ] [ Tanda Tangan & Cap ]), dan Baris Nama Direksi.')

    # --- BAGIAN 4 ---
    add_h1('BAGIAN 4: SPESIFIKASI & LOGIKA TOMBOL PILIHAN BAHASA (LANGUAGE SWITCHER)')
    add_p('Sistem web wajib menyediakan tombol pilihan bahasa yang responsif dan terintegrasi pada frontend dan backend:')

    add_h2('4.1 Tampilan Antarmuka Tombol Pemilihan Bahasa (UI Component)')
    add_callout(
        'Komponen UI ditempatkan pada form konfigurasi surat sebelum tombol generate/preview ditekan. Pengguna disajikan 3 opsi mutually-exclusive: 1) Bahasa Indonesia (Default), 2) British English, dan 3) Bilingual (Dwibahasa ID - EN-GB).',
        bold_prefix='KOMPONEN UI:'
    )

    ui_code = (
        '// Contoh Implementasi Komponen React (Tailwind CSS)\n'
        'export function LanguageSelector({ selectedLang, onSelectLang }) {\n'
        '  const languages = [\n'
        '    { id: "id", label: "Bahasa Indonesia", flag: "🇮🇩", desc: "Format Standar NBE (4 Halaman)" },\n'
        '    { id: "en", label: "British English", flag: "🇬🇧", desc: "UK Corporate Standard (4 Pages)" },\n'
        '    { id: "bi", label: "Bilingual (ID - EN)", flag: "🌐", desc: "Dual-Column Parallel (4 Pages)" }\n'
        '  ];\n'
        '  return (\n'
        '    <div className="flex flex-col gap-2 my-4 p-4 bg-slate-50 border rounded-lg">\n'
        '      <label className="text-sm font-bold text-slate-800">Pilih Bahasa Dokumen Output:</label>\n'
        '      <div className="grid grid-cols-3 gap-3">\n'
        '        {languages.map((l) => (\n'
        '          <button\n'
        '            key={l.id}\n'
        '            type="button"\n'
        '            onClick={() => onSelectLang(l.id)}\n'
        '            className={`flex flex-col items-center p-3 rounded-md border text-center transition ${\n'
        '              selectedLang === l.id\n'
        '                ? "bg-blue-900 text-white border-blue-900 shadow-sm"\n'
        '                : "bg-white text-slate-700 border-slate-300 hover:bg-slate-100"\n'
        '            }`}\n'
        '          >\n'
        '            <span className="text-xl mb-1">{l.flag}</span>\n'
        '            <span className="font-semibold text-sm">{l.label}</span>\n'
        '            <span className={`text-xs mt-1 ${selectedLang === l.id ? "text-blue-100" : "text-slate-500"}`}>{l.desc}</span>\n'
        '          </button>\n'
        '        ))}\n'
        '      </div>\n'
        '    </div>\n'
        '  );\n'
        '}'
    )
    add_code(ui_code)

    add_h2('4.2 Matriks Logika Pemrosesan Dokumen Berdasarkan Bahasa')
    t_lang = doc.add_table(rows=5, cols=4)
    t_lang.alignment = WD_TABLE_ALIGNMENT.CENTER
    lang_headers = ['Parameter', 'Opsi: Bahasa Indonesia', 'Opsi: British English', 'Opsi: Bilingual (ID - EN)']
    for c_i, h in enumerate(lang_headers):
        cell = t_lang.cell(0, c_i)
        cell.paragraphs[0].text = h
        set_cell_shading(cell, 'D9E1F2')
        set_cell_vcenter(cell)
        set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
        style_run(cell.paragraphs[0].runs[0], font_name='Arial', size_pt=8.5, bold=True, color_rgb=(0, 0, 0))

    lang_rows = [
        ('Judul Dokumen', 'SURAT KEPUTUSAN DIREKSI', "BOARD OF DIRECTORS' RESOLUTION", 'SURAT KEPUTUSAN DIREKSI / BOARD OF DIRECTORS\' RESOLUTION'),
        ('Konsiderans', 'Menimbang: & Mengingat:', 'Having Considered: & In View Of:', 'Menimbang / Having Considered & Mengingat / In View Of'),
        ('Diktum', 'MEMUTUSKAN, KESATU s.d. KELIMA', 'HEREBY RESOLVES, FIRST to FIFTH', 'Dual-Column Parallel (Kiri ID, Kanan EN)'),
        ('Tanda Tangan', 'Disetujui, Diverifikasi, Disiapkan', 'Approved, Verified, Prepared', 'Bilingual 3 Kolom Bersanding')
    ]
    for r_i, r_data in enumerate(lang_rows, start=1):
        row = t_lang.rows[r_i]
        set_row_cant_split(row)
        bg = 'FFFFFF' if r_i % 2 != 0 else 'F9FBFD'
        for c_i, val in enumerate(r_data):
            cell = row.cells[c_i]
            cell.paragraphs[0].text = val
            set_cell_shading(cell, bg)
            set_cell_vcenter(cell)
            set_cell_margins(cell, top=40, bottom=40, left=60, right=60)
            style_run(cell.paragraphs[0].runs[0], font_name='Arial', size_pt=8.0, bold=c_i == 0, color_rgb=(0x3B, 0x38, 0x38) if c_i > 0 else (0, 0, 0))

    # --- BAGIAN 5 ---
    add_h1('BAGIAN 5: KODE REFERENSI GENERATOR (PYTHON-DOCX BACKEND ENGINE)')
    add_p('Script generator berikut merupakan modul Python yang telah teruji menghasilkan ketiga varian dokumen di atas dengan ukuran tepat 4 halaman tanpa kesalahan:')

    snippet_python = (
        '# Cuplikan Engine Pemilihan Bahasa pada Backend (Python-Docx)\n'
        'def generate_nbe_letter(target_language="id", output_docx_path=None):\n'
        '    template_path = r"C:\\Users\\Mazmur Gusti Agung L\\Documents\\NBE\\Template Dokumen.docx"\n'
        '    doc = docx.Document(template_path)\n'
        '    \n'
        '    if target_language == "id":\n'
        '        # Bangun dokumen Bahasa Indonesia (4 Halaman)\n'
        '        build_indonesian_content(doc)\n'
        '    elif target_language == "en":\n'
        '        # Bangun dokumen British English (4 Halaman)\n'
        '        build_british_content(doc)\n'
        '    elif target_language == "bi":\n'
        '        # Bangun dokumen Dwibahasa Dual-Column Parallel (4 Halaman)\n'
        '        build_bilingual_content(doc)\n'
        '        \n'
        '    doc.save(output_docx_path)\n'
        '    # Konversi otomatis ke PDF melalui Word COM API\n'
        '    convert_docx_to_pdf(output_docx_path, output_docx_path.replace(".docx", ".pdf"))\n'
        '    return output_docx_path'
    )
    add_code(snippet_python)

    # Output file paths
    target_folder = r'C:\Users\Mazmur Gusti Agung L\Documents\NBE\Contoh surat'
    out_docx = os.path.join(target_folder, 'PANDUAN LENGKAP & PROMPT TERINTEGRASI - SISTEM WEB DOKUMEN NBE (ANTIGRAVITY IDE).docx')
    out_pdf = os.path.join(target_folder, 'PANDUAN LENGKAP & PROMPT TERINTEGRASI - SISTEM WEB DOKUMEN NBE (ANTIGRAVITY IDE).pdf')

    doc.save(out_docx)
    print(f"Generated Master Guide DOCX: {out_docx}")
    convert_docx_to_pdf(out_docx, out_pdf)
    return out_docx, out_pdf

if __name__ == '__main__':
    build_master_guide()
