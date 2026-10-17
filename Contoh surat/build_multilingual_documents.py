import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import win32com.client
import pymupdf

def set_cell_margins(cell, top=50, bottom=50, left=70, right=70):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    for s in tcPr.xpath('w:shd'):
        tcPr.remove(s)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_vcenter(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    for v in tcPr.xpath('w:vAlign'):
        tcPr.remove(v)
    vAlign = OxmlElement('w:vAlign')
    vAlign.set(qn('w:val'), 'center')
    tcPr.append(vAlign)

def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    cantSplit = OxmlElement('w:cantSplit')
    trPr.append(cantSplit)

def set_table_borders_none(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'  <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'  <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'  <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'  <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="D9D9D9"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def style_run(run, font_name='Arial', size_pt=10.0, bold=False, italic=False, color_rgb=(0x3B, 0x38, 0x38)):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = RGBColor(*color_rgb)

def convert_docx_to_pdf(docx_path, pdf_path):
    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    try:
        doc = word.Documents.Open(os.path.abspath(docx_path))
        doc.SaveAs(os.path.abspath(pdf_path), FileFormat=17) # wdFormatPDF = 17
        doc.Close()
    finally:
        word.Quit()
    print(f"Exported PDF successfully: {pdf_path}")

# ==============================================================================
# 1. BUILD BRITISH ENGLISH DOCUMENT
# ==============================================================================
def build_british_english():
    template_path = r'C:\Users\Mazmur Gusti Agung L\Documents\NBE\Template Dokumen.docx'
    target_folder = r'C:\Users\Mazmur Gusti Agung L\Documents\NBE\Contoh surat'
    out_docx = os.path.join(target_folder, '0002-KB-GNR-CEO-X-2026 - Board of Directors Resolution (British English).docx')
    out_pdf = os.path.join(target_folder, '0002-KB-GNR-CEO-X-2026 - Board of Directors Resolution (British English).pdf')

    doc = docx.Document(template_path)

    # --- KOP HEADER (Page 2 onwards) ---
    ht = doc.sections[0].header.tables[0]
    cell_div = ht.cell(0, 1)
    cell_div.paragraphs[0].text = ''
    r = cell_div.paragraphs[0].add_run('GENERAL & MANAGEMENT (GNR)')
    style_run(r, font_name='Arial', size_pt=9.5, bold=True, color_rgb=(0, 0, 0))
    cell_div.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    cell_cat = ht.cell(1, 1)
    cell_cat.paragraphs[0].text = ''
    r = cell_cat.paragraphs[0].add_run('COMPANY POLICY (KB)')
    style_run(r, font_name='Arial', size_pt=9.5, bold=True, color_rgb=(0, 0, 0))
    cell_cat.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    cell_title = ht.cell(2, 1)
    cell_title.paragraphs[0].text = ''
    r = cell_title.paragraphs[0].add_run("BOARD OF DIRECTORS' RESOLUTION: DOCUMENT CONTROL SYSTEM & DCC TEAM")
    style_run(r, font_name='Arial', size_pt=9.0, bold=True, color_rgb=(0, 0, 0))
    cell_title.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    c_rev = ht.cell(3, 1)
    c_rev.paragraphs[0].text = ''
    r = c_rev.paragraphs[0].add_run('Revision: 0')
    style_run(r, font_name='Arial', size_pt=8.0, bold=False, color_rgb=(0, 0, 0))
    c_rev.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    c_code = ht.cell(3, 2)
    c_code.paragraphs[0].text = ''
    r = c_code.paragraphs[0].add_run('0002-KB-GNR-CEO-X-2026')
    style_run(r, font_name='Arial', size_pt=8.0, bold=True, color_rgb=(0, 0, 0))
    c_code.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    c_date = ht.cell(3, 3)
    c_date.paragraphs[0].text = ''
    r = c_date.paragraphs[0].add_run('Date : 17 October 2026')
    style_run(r, font_name='Arial', size_pt=8.0, bold=False, color_rgb=(0, 0, 0))
    c_date.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    for r_i in range(len(ht.rows)):
        for c_i in range(len(ht.columns)):
            cell = ht.cell(r_i, c_i)
            set_cell_vcenter(cell)
            set_cell_margins(cell, top=50, bottom=50, left=80, right=80)

    # --- COVER PAGE (PAGE 1) ---
    p19 = doc.paragraphs[19]
    p19.text = ''
    r = p19.add_run('PT NIRWANA BHUMI ENERGY')
    style_run(r, font_name='Arial', size_pt=20.0, bold=True, color_rgb=(0, 0, 0))
    p19.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p19.paragraph_format.space_before = Pt(12)
    p19.paragraph_format.space_after = Pt(6)

    p21 = doc.paragraphs[21]
    p21.text = ''
    r1 = p21.add_run("BOARD OF DIRECTORS' RESOLUTION\n")
    style_run(r1, font_name='Arial', size_pt=16.0, bold=True, color_rgb=(0, 0, 0))
    r2 = p21.add_run('CONCERNING THE ESTABLISHMENT OF THE DOCUMENT CONTROL MANAGEMENT SYSTEM\nAND THE APPOINTMENT OF THE DOCUMENT CONTROL CENTRE (DCC) TEAM\n\n')
    style_run(r2, font_name='Arial', size_pt=11.5, bold=True, color_rgb=(0, 0, 0))
    r3 = p21.add_run('Document Number: 0002-KB-GNR-CEO-X-2026')
    style_run(r3, font_name='Arial', size_pt=11.0, bold=True, color_rgb=(0, 51, 102))
    p21.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p21.paragraph_format.space_after = Pt(6)

    for p_idx in range(22, 27):
        p = doc.paragraphs[p_idx]
        p.text = ''
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(13)

    for p_idx in range(27, 32):
        p = doc.paragraphs[p_idx]
        p.text = ''
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(2)

    # Approval block at bottom-right
    p32 = doc.paragraphs[32]
    p32.text = ''
    p32.paragraph_format.left_indent = Inches(3.65)
    p32.paragraph_format.space_before = Pt(0)
    p32.paragraph_format.space_after = Pt(1)
    p32.paragraph_format.line_spacing = 1.05
    r_l1 = p32.add_run('Approved by   : ')
    style_run(r_l1, font_name='Arial', size_pt=9.0, bold=True, color_rgb=(0, 0, 0))
    r_v1 = p32.add_run('Bima Satryo Hartanto\n(President Director / CEO)')
    style_run(r_v1, font_name='Arial', size_pt=9.0, bold=False, color_rgb=(0, 0, 0))

    p33 = doc.paragraphs[33]
    p33.text = ''
    p33.paragraph_format.left_indent = Inches(3.65)
    p33.paragraph_format.space_before = Pt(3)
    p33.paragraph_format.space_after = Pt(1)
    p33.paragraph_format.line_spacing = 1.05
    r_l2 = p33.add_run('Verified by    : ')
    style_run(r_l2, font_name='Arial', size_pt=9.0, bold=True, color_rgb=(0, 0, 0))
    r_v2 = p33.add_run('Dennis Raditya (Director of\nOperations & Technology / CTO)')
    style_run(r_v2, font_name='Arial', size_pt=9.0, bold=False, color_rgb=(0, 0, 0))

    p34 = doc.paragraphs[34]
    p34.text = ''
    p34.paragraph_format.left_indent = Inches(3.65)
    p34.paragraph_format.space_before = Pt(3)
    p34.paragraph_format.space_after = Pt(0)
    p34.paragraph_format.line_spacing = 1.05
    r_l3 = p34.add_run('Prepared by   : ')
    style_run(r_l3, font_name='Arial', size_pt=9.0, bold=True, color_rgb=(0, 0, 0))
    r_v3 = p34.add_run('Mazmur Gusti Agung Larosa\n(Director of Finance & Engineering / DCC Lead)')
    style_run(r_v3, font_name='Arial', size_pt=9.0, bold=False, color_rgb=(0, 0, 0))

    # Page Break to Page 2
    p35 = doc.paragraphs[35]
    p35.text = ''
    p35.add_run().add_break(WD_BREAK.PAGE)

    # --- PAGE 2: REVISION INDEX ---
    p36 = doc.paragraphs[36]
    p36.text = ''
    r_rev = p36.add_run('REVISION INDEX')
    style_run(r_rev, font_name='Arial', size_pt=12.0, bold=True, color_rgb=(0, 0, 0))
    p36.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p36.paragraph_format.space_before = Pt(18)
    p36.paragraph_format.space_after = Pt(12)

    t_rev = doc.tables[0]
    set_row_cant_split(t_rev.rows[0])
    set_row_cant_split(t_rev.rows[1])

    headers_en = ['Rev.', 'Date', 'Description', 'Prepared By']
    for c_i, h in enumerate(headers_en):
        cell = t_rev.rows[0].cells[c_i]
        cell.paragraphs[0].text = h
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(cell, 'D9E1F2')
        set_cell_vcenter(cell)
        set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
        for r in cell.paragraphs[0].runs:
            style_run(r, font_name='Arial', size_pt=9.5, bold=True, color_rgb=(0, 0, 0))

    row1 = t_rev.rows[1]
    row1.cells[0].paragraphs[0].text = '0'
    row1.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    row1.cells[1].paragraphs[0].text = '17 October 2026'
    row1.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    row1.cells[2].paragraphs[0].text = 'First Edition – Establishment of the Document Control Management System and Appointment of the Document Control Centre (DCC) Team of PT Nirwana Bhumi Energy'
    row1.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT

    row1.cells[3].paragraphs[0].text = 'Mazmur Gusti Agung Larosa'
    row1.cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT

    for cell in row1.cells:
        set_cell_shading(cell, 'FFFFFF')
        set_cell_vcenter(cell)
        set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
        for p in cell.paragraphs:
            for r in p.runs:
                style_run(r, font_name='Arial', size_pt=9.0, bold=False, color_rgb=(0x3B, 0x38, 0x38))

    if len(t_rev.rows) > 2:
        tr = t_rev.rows[2]._tr
        tr.getparent().remove(tr)

    doc.paragraphs[37].text = ''
    doc.paragraphs[38].text = ''

    # Page Break to Page 3
    p_pb_h2 = doc.add_paragraph()
    p_pb_h2.add_run().add_break(WD_BREAK.PAGE)

    # --- PAGE 3: RESOLUTION BODY ---
    def add_p(text='', bold=False, italic=False, size_pt=9.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=2, left_indent=0, color_rgb=(0x3B, 0x38, 0x38), keep_with_next=False):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.10
        p.paragraph_format.keep_with_next = keep_with_next
        if left_indent > 0:
            p.paragraph_format.left_indent = Inches(left_indent)
        if text:
            r = p.add_run(text)
            style_run(r, font_name='Arial', size_pt=size_pt, bold=bold, italic=italic, color_rgb=color_rgb)
        return p

    add_p("BOARD OF DIRECTORS' RESOLUTION OF PT NIRWANA BHUMI ENERGY", bold=True, size_pt=11.0, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=1, color_rgb=(0, 0, 0), keep_with_next=True)
    add_p('NUMBER: 0002-KB-GNR-CEO-X-2026', bold=True, size_pt=10.0, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2, color_rgb=(0, 51, 102), keep_with_next=True)
    add_p('CONCERNING', bold=True, size_pt=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1, color_rgb=(0, 0, 0), keep_with_next=True)
    add_p('THE ESTABLISHMENT OF THE DOCUMENT CONTROL MANAGEMENT SYSTEM AND THE APPOINTMENT OF THE DOCUMENT CONTROL CENTRE (DCC) TEAM', bold=True, size_pt=10.0, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8, color_rgb=(0, 0, 0), keep_with_next=True)

    add_p('THE BOARD OF DIRECTORS OF PT NIRWANA BHUMI ENERGY,', bold=True, size_pt=9.5, space_before=2, space_after=4, color_rgb=(0, 0, 0), keep_with_next=True)

    add_p('Having Considered:', bold=True, size_pt=9.5, space_before=2, space_after=1, color_rgb=(0, 0, 0), keep_with_next=True)
    add_p('a.  that in order to achieve standardisation, accountability, and traceability of all corporate policy documents, standard operating procedures, work instructions, and technical forms within PT Nirwana Bhumi Energy, a standardised and integrated document governance and control system is required;', left_indent=0.4, space_after=2)
    add_p('b.  that to ensure comprehensive compliance with Integrated Management System (IMS) standards based on ISO 9001:2015, ISO 14001:2015, ISO 45001:2018, and Contractor Safety Management System (CSMS) qualifications of SKK Migas and PT PLN (Persero), the appointment of personnel responsible for document management is requisite;', left_indent=0.4, space_after=2)
    add_p("c.  that based on the considerations referred to in paragraphs a and b above, it is deemed necessary to enact a Board of Directors' Resolution concerning the Establishment of the Document Control Management System and the Appointment of the Document Control Centre (DCC) Team of PT Nirwana Bhumi Energy.", left_indent=0.4, space_after=4)

    add_p('In View Of:', bold=True, size_pt=9.5, space_before=2, space_after=1, color_rgb=(0, 0, 0), keep_with_next=True)
    add_p('1.  The Deed of Establishment of PT Nirwana Bhumi Energy and prevailing statutory regulations governing limited liability companies;', left_indent=0.4, space_after=1.5)
    add_p('2.  International Standard ISO 9001:2015 Clause 7.5 regarding Documented Information;', left_indent=0.4, space_after=1.5)
    add_p('3.  Standard Guidelines for the Contractor Safety Management System (CSMS) of SKK Migas and PT PLN (Persero);', left_indent=0.4, space_after=1.5)
    add_p('4.  Standard Operating Procedure No. 0001-SOP-GNR-CEO-IX-2026 (Documents and Records Control Procedure);', left_indent=0.4, space_after=1.5)
    add_p('5.  Integrated Management Policy No. 0001-KB-GNR-CEO-X-2026.', left_indent=0.4, space_after=4)

    add_p('HEREBY RESOLVES:', bold=True, size_pt=10.0, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=4, color_rgb=(0, 0, 0), keep_with_next=True)

    add_p('To Enact:', bold=True, size_pt=9.5, space_after=2, color_rgb=(0, 0, 0), keep_with_next=True)
    add_p("BOARD OF DIRECTORS' RESOLUTION OF PT NIRWANA BHUMI ENERGY CONCERNING THE ESTABLISHMENT OF THE DOCUMENT CONTROL MANAGEMENT SYSTEM AND THE APPOINTMENT OF THE DOCUMENT CONTROL CENTRE (DCC) TEAM.", bold=True, left_indent=0.4, space_after=4, color_rgb=(0, 0, 0), keep_with_next=True)

    add_p('FIRST :', bold=True, size_pt=9.5, space_before=3, space_after=1, color_rgb=(0, 0, 0), keep_with_next=True)
    add_p('To officially enact the Standard Format and 6-Segment Document Numbering System across all operational areas and the head office of PT Nirwana Bhumi Energy, with the following mandatory structure:', left_indent=0.4, space_after=2)
    add_p('[4DigitSequentialNo] - [DocCategory] - [Division] - [Authoriser] - [RomanMonth] - [Year]', bold=True, left_indent=0.6, space_after=2, color_rgb=(0, 51, 102))
    add_p('This numbering provision is strictly unique under the rigorous principle: ONE CODE FOR ONE DOCUMENT (1 Document Code = 1 Valid Document). Any duplication of document numbers or reuse of the same number for different documents is strictly prohibited by the system.', left_indent=0.4, space_after=4)

    # Page Break to Page 4
    p_pb_h3 = doc.add_paragraph()
    p_pb_h3.add_run().add_break(WD_BREAK.PAGE)

    # --- PAGE 4: DCC TEAM, CLAUSES, SIGNATURES ---
    add_p('SECOND :', bold=True, size_pt=9.5, space_before=2, space_after=2, color_rgb=(0, 0, 0), keep_with_next=True)
    add_p('To appoint and designate the following personnel as the Document Control Centre (DCC) Team of PT Nirwana Bhumi Energy:', left_indent=0.4, space_after=3, keep_with_next=True)

    table_team = doc.add_table(rows=4, cols=4)
    table_team.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ['No.', 'Personnel Name', 'Structural Position', 'DCC Role & Responsibilities']
    col_widths = [Inches(0.4), Inches(1.8), Inches(1.7), Inches(2.5)]

    for c_i, name in enumerate(headers):
        cell = table_team.cell(0, c_i)
        cell.width = col_widths[c_i]
        cell.paragraphs[0].text = name
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(cell, 'D9E1F2')
        set_cell_vcenter(cell)
        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
        for r in cell.paragraphs[0].runs:
            style_run(r, font_name='Arial', size_pt=8.5, bold=True, color_rgb=(0, 0, 0))

    team_data = [
        ('1.', 'Mazmur Gusti Agung Larosa', 'Director of Finance & Engineering', "Lead Document Controller & System Administrator:\nValidates drafts, registers unique codes, executes the 'Fix Code' command, and synchronises corporate Google Drive archives."),
        ('2.', 'Dennis Raditya', 'Director of Operations & Technology', 'Technical Verifier & QA Approver:\nReviews technical format conformity, sub-pixel layout alignment, and compliance with IMS & CSMS quality standards.'),
        ('3.', 'Abdul Qodir Al Jailani', 'Commissioner / IMS Lead', 'Independent Compliance Overseer:\nSupervises periodic audits of controlled documents, past numbering (backdating), and archival retention integrity.')
    ]

    for r_i, data in enumerate(team_data, start=1):
        row = table_team.rows[r_i]
        set_row_cant_split(row)
        bg = 'FFFFFF' if r_i % 2 != 0 else 'F9FBFD'
        for c_i, val in enumerate(data):
            cell = row.cells[c_i]
            cell.width = col_widths[c_i]
            cell.paragraphs[0].text = val
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if c_i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_shading(cell, bg)
            set_cell_vcenter(cell)
            set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
            for r in cell.paragraphs[0].runs:
                style_run(r, font_name='Arial', size_pt=8.0, bold=False, color_rgb=(0x3B, 0x38, 0x38))

    add_p(space_after=3)

    add_p('THIRD :', bold=True, size_pt=9.5, space_before=3, space_after=1, color_rgb=(0, 0, 0), keep_with_next=True)
    add_p("To govern the digital storage and registration of corporate documents via the 'Fix Code' button on the web system, whereby upon code verification and locking, the system automatically synchronises and archives the document into corporate Google Drive (Root Folder ID: 1ny_1VhXfSaNrj_XdoOsC7V7BO-z0K8eF) within its respective divisional classification subfolder.", left_indent=0.4, space_after=3)

    add_p('FOURTH :', bold=True, size_pt=9.5, space_before=3, space_after=1, color_rgb=(0, 0, 0), keep_with_next=True)
    add_p('To establish the revision workflow for controlled documents, whereby any subsequent amendment to an issued document must append a suffix (-rev1, -rev2, etc.) to the digital file, record the specific rationale in detail within the Revision Index table, and promptly withdraw and destroy obsolete hard copies and digital versions from active circulation.', left_indent=0.4, space_after=3)

    add_p('FIFTH :', bold=True, size_pt=9.5, space_before=3, space_after=1, color_rgb=(0, 0, 0), keep_with_next=True)
    add_p('This Resolution shall enter into force effectively from the date of its enactment. Should any discrepancies or regulatory adjustments within the oil & gas or power industry arise in the future, due evaluation and rectification shall be executed accordingly.', left_indent=0.4, space_after=6)

    add_p('Enacted in : Semarang / Tuban\nOn date    : 17 October 2026', bold=False, size_pt=9.0, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=3, keep_with_next=True)
    add_p('BOARD OF DIRECTORS OF PT NIRWANA BHUMI ENERGY', bold=True, size_pt=9.5, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4, color_rgb=(0, 0, 0), keep_with_next=True)

    t_sig = doc.add_table(rows=3, cols=3)
    t_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_row_cant_split(t_sig.rows[0])
    set_row_cant_split(t_sig.rows[1])
    set_row_cant_split(t_sig.rows[2])

    sig_roles = [
        'Approved by:\nPresident Director / CEO',
        'Verified by:\nDirector of Operations & Tech (CTO)',
        'Prepared & Proposed by:\nDirector of Finance & Eng (CFO / DCC)'
    ]
    sig_names = [
        'BIMA SATRYO HARTANTO\nPresident Director',
        'DENNIS RADITYA\nDirector of Operations',
        'MAZMUR GUSTI AGUNG LAROSA\nDirector of Finance & Engineering'
    ]

    for c_i, role in enumerate(sig_roles):
        cell = t_sig.cell(0, c_i)
        cell.paragraphs[0].text = role
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_vcenter(cell)
        set_cell_margins(cell, top=40, bottom=40, left=60, right=60)
        for r in cell.paragraphs[0].runs:
            style_run(r, font_name='Arial', size_pt=8.0, bold=True, color_rgb=(0, 0, 0))

    for c_i in range(3):
        cell = t_sig.cell(1, c_i)
        cell.paragraphs[0].text = '\n[ Signature & Seal ]\n' if c_i != 0 else '\n[ Stamp Duty Rp 10,000 ]\n[ Signature & Seal ]\n'
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_vcenter(cell)
        set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
        for r in cell.paragraphs[0].runs:
            style_run(r, font_name='Arial', size_pt=7.5, italic=True, color_rgb=(0x88, 0x88, 0x88))

    for c_i, name in enumerate(sig_names):
        cell = t_sig.cell(2, c_i)
        cell.paragraphs[0].text = name
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_vcenter(cell)
        set_cell_margins(cell, top=40, bottom=40, left=60, right=60)
        for r in cell.paragraphs[0].runs:
            style_run(r, font_name='Arial', size_pt=8.0, bold=True, color_rgb=(0, 0, 0))

    doc.save(out_docx)
    print(f"Generated DOCX (British English): {out_docx}")
    convert_docx_to_pdf(out_docx, out_pdf)
    return out_docx, out_pdf

# ==============================================================================
# 2. BUILD BILINGUAL (BAHASA INDONESIA & BRITISH ENGLISH) DOCUMENT
# ==============================================================================
def build_bilingual():
    template_path = r'C:\Users\Mazmur Gusti Agung L\Documents\NBE\Template Dokumen.docx'
    target_folder = r'C:\Users\Mazmur Gusti Agung L\Documents\NBE\Contoh surat'
    out_docx = os.path.join(target_folder, '0002-KB-GNR-CEO-X-2026 - Surat Keputusan Direksi (Bilingual ID-EN).docx')
    out_pdf = os.path.join(target_folder, '0002-KB-GNR-CEO-X-2026 - Surat Keputusan Direksi (Bilingual ID-EN).pdf')

    doc = docx.Document(template_path)

    # --- KOP HEADER (Page 2 onwards) ---
    ht = doc.sections[0].header.tables[0]
    cell_div = ht.cell(0, 1)
    cell_div.paragraphs[0].text = ''
    r = cell_div.paragraphs[0].add_run('GENERAL & MANAGEMENT (GNR)')
    style_run(r, font_name='Arial', size_pt=9.5, bold=True, color_rgb=(0, 0, 0))
    cell_div.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    cell_cat = ht.cell(1, 1)
    cell_cat.paragraphs[0].text = ''
    r = cell_cat.paragraphs[0].add_run('KEBIJAKAN PERUSAHAAN / COMPANY POLICY (KB)')
    style_run(r, font_name='Arial', size_pt=9.0, bold=True, color_rgb=(0, 0, 0))
    cell_cat.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    cell_title = ht.cell(2, 1)
    cell_title.paragraphs[0].text = ''
    r = cell_title.paragraphs[0].add_run('SK DIREKSI: SISTEM PENGENDALIAN DOKUMEN & TIM DCC\nBOD RESOLUTION: DOCUMENT CONTROL SYSTEM & DCC TEAM')
    style_run(r, font_name='Arial', size_pt=8.5, bold=True, color_rgb=(0, 0, 0))
    cell_title.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    c_rev = ht.cell(3, 1)
    c_rev.paragraphs[0].text = ''
    r = c_rev.paragraphs[0].add_run('Revision: 0')
    style_run(r, font_name='Arial', size_pt=8.0, bold=False, color_rgb=(0, 0, 0))
    c_rev.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    c_code = ht.cell(3, 2)
    c_code.paragraphs[0].text = ''
    r = c_code.paragraphs[0].add_run('0002-KB-GNR-CEO-X-2026')
    style_run(r, font_name='Arial', size_pt=8.0, bold=True, color_rgb=(0, 0, 0))
    c_code.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    c_date = ht.cell(3, 3)
    c_date.paragraphs[0].text = ''
    r = c_date.paragraphs[0].add_run('Date : 17 Oktober 2026')
    style_run(r, font_name='Arial', size_pt=8.0, bold=False, color_rgb=(0, 0, 0))
    c_date.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    for r_i in range(len(ht.rows)):
        for c_i in range(len(ht.columns)):
            cell = ht.cell(r_i, c_i)
            set_cell_vcenter(cell)
            set_cell_margins(cell, top=50, bottom=50, left=80, right=80)

    # --- COVER PAGE (PAGE 1) ---
    p19 = doc.paragraphs[19]
    p19.text = ''
    r = p19.add_run('PT NIRWANA BHUMI ENERGY')
    style_run(r, font_name='Arial', size_pt=20.0, bold=True, color_rgb=(0, 0, 0))
    p19.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p19.paragraph_format.space_before = Pt(12)
    p19.paragraph_format.space_after = Pt(4)

    p21 = doc.paragraphs[21]
    p21.text = ''
    r1 = p21.add_run('SURAT KEPUTUSAN DIREKSI\n')
    style_run(r1, font_name='Arial', size_pt=15.0, bold=True, color_rgb=(0, 0, 0))
    r1_en = p21.add_run("BOARD OF DIRECTORS' RESOLUTION\n\n")
    style_run(r1_en, font_name='Arial', size_pt=13.0, bold=True, italic=True, color_rgb=(0x44, 0x44, 0x44))
    
    r2 = p21.add_run('TENTANG PENETAPAN SISTEM MANAJEMEN PENGENDALIAN DOKUMEN\nDAN PENUNJUKAN TIM DOCUMENT CONTROL CENTER (DCC)\n')
    style_run(r2, font_name='Arial', size_pt=11.0, bold=True, color_rgb=(0, 0, 0))
    r2_en = p21.add_run('CONCERNING THE ESTABLISHMENT OF THE DOCUMENT CONTROL MANAGEMENT SYSTEM\nAND THE APPOINTMENT OF THE DOCUMENT CONTROL CENTRE (DCC) TEAM\n\n')
    style_run(r2_en, font_name='Arial', size_pt=10.0, bold=True, italic=True, color_rgb=(0x44, 0x44, 0x44))

    r3 = p21.add_run('Nomor Dokumen / Document Number: 0002-KB-GNR-CEO-X-2026')
    style_run(r3, font_name='Arial', size_pt=10.5, bold=True, color_rgb=(0, 51, 102))
    p21.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p21.paragraph_format.space_after = Pt(4)

    # Space down
    for p_idx in range(22, 26):
        p = doc.paragraphs[p_idx]
        p.text = ''
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(11)

    for p_idx in range(26, 32):
        p = doc.paragraphs[p_idx]
        p.text = ''
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(2)

    # Approval block at bottom right
    p32 = doc.paragraphs[32]
    p32.text = ''
    p32.paragraph_format.left_indent = Inches(3.4)
    p32.paragraph_format.space_before = Pt(0)
    p32.paragraph_format.space_after = Pt(1)
    p32.paragraph_format.line_spacing = 1.05
    r_l1 = p32.add_run('Approved by / Disetujui  : ')
    style_run(r_l1, font_name='Arial', size_pt=8.5, bold=True, color_rgb=(0, 0, 0))
    r_v1 = p32.add_run('Bima Satryo Hartanto\n(President Director / CEO)')
    style_run(r_v1, font_name='Arial', size_pt=8.5, bold=False, color_rgb=(0, 0, 0))

    p33 = doc.paragraphs[33]
    p33.text = ''
    p33.paragraph_format.left_indent = Inches(3.4)
    p33.paragraph_format.space_before = Pt(2)
    p33.paragraph_format.space_after = Pt(1)
    p33.paragraph_format.line_spacing = 1.05
    r_l2 = p33.add_run('Verified by / Diverifikasi : ')
    style_run(r_l2, font_name='Arial', size_pt=8.5, bold=True, color_rgb=(0, 0, 0))
    r_v2 = p33.add_run('Dennis Raditya (Director of\nOperations & Technology / CTO)')
    style_run(r_v2, font_name='Arial', size_pt=8.5, bold=False, color_rgb=(0, 0, 0))

    p34 = doc.paragraphs[34]
    p34.text = ''
    p34.paragraph_format.left_indent = Inches(3.4)
    p34.paragraph_format.space_before = Pt(2)
    p34.paragraph_format.space_after = Pt(0)
    p34.paragraph_format.line_spacing = 1.05
    r_l3 = p34.add_run('Prepared by / Disiapkan   : ')
    style_run(r_l3, font_name='Arial', size_pt=8.5, bold=True, color_rgb=(0, 0, 0))
    r_v3 = p34.add_run('Mazmur Gusti Agung Larosa\n(Director of Finance & Eng / DCC Lead)')
    style_run(r_v3, font_name='Arial', size_pt=8.5, bold=False, color_rgb=(0, 0, 0))

    # Page Break to Page 2
    p35 = doc.paragraphs[35]
    p35.text = ''
    p35.add_run().add_break(WD_BREAK.PAGE)

    # --- PAGE 2: REVISION INDEX ---
    p36 = doc.paragraphs[36]
    p36.text = ''
    r_rev = p36.add_run('REVISION INDEX / INDEKS REVISI')
    style_run(r_rev, font_name='Arial', size_pt=12.0, bold=True, color_rgb=(0, 0, 0))
    p36.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p36.paragraph_format.space_before = Pt(18)
    p36.paragraph_format.space_after = Pt(12)

    t_rev = doc.tables[0]
    set_row_cant_split(t_rev.rows[0])
    set_row_cant_split(t_rev.rows[1])

    headers_bi = ['Rev.', 'Date / Tanggal', 'Description / Deskripsi Perubahan', 'Prepared By / Disusun Oleh']
    for c_i, h in enumerate(headers_bi):
        cell = t_rev.rows[0].cells[c_i]
        cell.paragraphs[0].text = h
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(cell, 'D9E1F2')
        set_cell_vcenter(cell)
        set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
        for r in cell.paragraphs[0].runs:
            style_run(r, font_name='Arial', size_pt=9.0, bold=True, color_rgb=(0, 0, 0))

    row1 = t_rev.rows[1]
    row1.cells[0].paragraphs[0].text = '0'
    row1.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    row1.cells[1].paragraphs[0].text = '17 Oktober 2026 /\n17 October 2026'
    row1.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    row1.cells[2].paragraphs[0].text = 'First Edition / Edisi Perdana – Penetapan Sistem Manajemen Pengendalian Dokumen dan Penunjukan Tim Document Control Center (DCC) PT Nirwana Bhumi Energy / Establishment of Document Control Management System and Appointment of Document Control Centre (DCC) Team.'
    row1.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT

    row1.cells[3].paragraphs[0].text = 'Mazmur Gusti Agung Larosa'
    row1.cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT

    for cell in row1.cells:
        set_cell_shading(cell, 'FFFFFF')
        set_cell_vcenter(cell)
        set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
        for p in cell.paragraphs:
            for r in p.runs:
                style_run(r, font_name='Arial', size_pt=8.5, bold=False, color_rgb=(0x3B, 0x38, 0x38))

    if len(t_rev.rows) > 2:
        tr = t_rev.rows[2]._tr
        tr.getparent().remove(tr)

    doc.paragraphs[37].text = ''
    doc.paragraphs[38].text = ''

    # Page Break to Page 3
    p_pb_h2 = doc.add_paragraph()
    p_pb_h2.add_run().add_break(WD_BREAK.PAGE)

    # --- PAGE 3: BILINGUAL RESOLUTION BODY (DUAL-COLUMN FORMAT) ---
    def add_p_header(text='', bold=False, italic=False, size_pt=10.0, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=2, color_rgb=(0, 0, 0), keep_with_next=True):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.keep_with_next = keep_with_next
        if text:
            r = p.add_run(text)
            style_run(r, font_name='Arial', size_pt=size_pt, bold=bold, italic=italic, color_rgb=color_rgb)
        return p

    add_p_header('SURAT KEPUTUSAN DIREKSI PT NIRWANA BHUMI ENERGY\nBOARD OF DIRECTORS\' RESOLUTION OF PT NIRWANA BHUMI ENERGY', bold=True, size_pt=10.0, space_before=1, space_after=1, color_rgb=(0, 0, 0))
    add_p_header('NOMOR / NUMBER: 0002-KB-GNR-CEO-X-2026', bold=True, size_pt=9.0, space_after=1, color_rgb=(0, 51, 102))
    add_p_header('TENTANG PENETAPAN SISTEM MANAJEMEN PENGENDALIAN DOKUMEN DAN PENUNJUKAN TIM DCC\nCONCERNING THE ESTABLISHMENT OF DOCUMENT CONTROL SYSTEM AND DCC TEAM', bold=True, size_pt=9.0, space_after=4, color_rgb=(0, 0, 0))

    # Helper function for dual-column table row with LEFT ALIGNMENT for crisp typography
    def add_dual_row(table, text_id, text_en, is_title=False, is_bold=False, bg_color=None, font_size=7.8):
        row = table.add_row()
        set_row_cant_split(row)
        cell_id = row.cells[0]
        cell_en = row.cells[1]
        cell_id.width = Inches(3.22)
        cell_en.width = Inches(3.22)

        # Process Indonesian lines
        cell_id.paragraphs[0].text = ''
        lines_id = text_id.split('\n') if isinstance(text_id, str) else text_id
        for idx, line in enumerate(lines_id):
            p = cell_id.paragraphs[0] if idx == 0 else cell_id.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1.0 if len(lines_id) > 1 else 0.5)
            p.paragraph_format.line_spacing = 1.02
            r = p.add_run(line)
            style_run(r, font_name='Arial', size_pt=font_size if not is_title else font_size + 0.5, bold=is_bold, italic=False, color_rgb=(0, 0, 0) if is_bold else (0x3B, 0x38, 0x38))

        # Process English lines
        cell_en.paragraphs[0].text = ''
        lines_en = text_en.split('\n') if isinstance(text_en, str) else text_en
        for idx, line in enumerate(lines_en):
            p = cell_en.paragraphs[0] if idx == 0 else cell_en.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1.0 if len(lines_en) > 1 else 0.5)
            p.paragraph_format.line_spacing = 1.02
            r = p.add_run(line)
            style_run(r, font_name='Arial', size_pt=font_size if not is_title else font_size + 0.5, bold=is_bold, italic=True if not is_bold else False, color_rgb=(0, 0, 0) if is_bold else (0x33, 0x33, 0x33))

        for c in (cell_id, cell_en):
            set_cell_vcenter(c)
            set_cell_margins(c, top=18, bottom=18, left=35, right=35)
            if bg_color:
                set_cell_shading(c, bg_color)

    # Table for Recitals & Opening (Page 3)
    t_recitals = doc.add_table(rows=0, cols=2)
    t_recitals.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_none(t_recitals)

    # Opening
    add_dual_row(t_recitals, 'DIREKSI PT NIRWANA BHUMI ENERGY,', 'THE BOARD OF DIRECTORS OF PT NIRWANA BHUMI ENERGY,', is_title=True, is_bold=True)
    
    # Menimbang
    add_dual_row(t_recitals, 'Menimbang:', 'Having Considered:', is_title=True, is_bold=True)
    add_dual_row(t_recitals, 
                 'a. bahwa dalam rangka mewujudkan standarisasi, akuntabilitas, serta keterlacakan seluruh dokumen kebijakan, SOP, IK, dan formulir teknis di lingkungan PT Nirwana Bhumi Energy, diperlukan sistem tata kelola dan kontrol dokumen yang baku dan terintegrasi;',
                 'a. that in order to achieve standardisation, accountability, and traceability of all policy documents, SOPs, WIs, and technical forms within PT Nirwana Bhumi Energy, a standardised and integrated document governance and control system is required;')
    add_dual_row(t_recitals,
                 'b. bahwa guna menjamin kepatuhan menyeluruh terhadap standar Sistem Manajemen Terintegrasi (IMS: ISO 9001, ISO 14001, ISO 45001) serta kualifikasi CSMS SKK Migas dan PLN, diperlukan penunjukan personel penanggung jawab pengelolaan dokumen;',
                 'b. that to ensure comprehensive compliance with Integrated Management System standards (IMS: ISO 9001, ISO 14001, ISO 45001) and CSMS qualifications of SKK Migas and PLN, the appointment of personnel responsible for document management is requisite;')
    add_dual_row(t_recitals,
                 'c. bahwa berdasarkan pertimbangan sebagaimana dimaksud dalam huruf a dan b di atas, dipandang perlu untuk menetapkan Surat Keputusan Direksi tentang Penetapan Sistem Manajemen Pengendalian Dokumen dan Penunjukan Tim Document Control Center (DCC).',
                 "c. that based on the considerations referred to in paragraphs a and b above, it is deemed necessary to enact a Board of Directors' Resolution concerning the Establishment of the Document Control Management System and Appointment of the DCC Team.")

    # Mengingat
    mengingat_id = (
        '1. Akta Pendirian PT Nirwana Bhumi Energy serta peraturan perseroan berlaku;\n'
        '2. Standar ISO 9001:2015 Klausul 7.5 (Documented Information);\n'
        '3. Pedoman CSMS SKK Migas dan PT PLN (Persero);\n'
        '4. Prosedur No. 0001-SOP-GNR-CEO-IX-2026 (Documents & Records Control);\n'
        '5. Kebijakan Terintegrasi No. 0001-KB-GNR-CEO-X-2026.'
    )
    mengingat_en = (
        '1. The Deed of Establishment of PT Nirwana Bhumi Energy & applicable company laws;\n'
        '2. International Standard ISO 9001:2015 Clause 7.5 (Documented Information);\n'
        '3. CSMS Guidelines of SKK Migas and PT PLN (Persero);\n'
        '4. Procedure No. 0001-SOP-GNR-CEO-IX-2026 (Documents & Records Control);\n'
        '5. Integrated Policy No. 0001-KB-GNR-CEO-X-2026.'
    )
    add_dual_row(t_recitals, 'Mengingat:', 'In View Of:', is_title=True, is_bold=True)
    add_dual_row(t_recitals, mengingat_id, mengingat_en)

    # MEMUTUSKAN
    add_dual_row(t_recitals, 'MEMUTUSKAN:', 'HEREBY RESOLVES:', is_title=True, is_bold=True, bg_color='F2F4F8')
    add_dual_row(t_recitals, 'Menetapkan:', 'To Enact:', is_title=True, is_bold=True)
    add_dual_row(t_recitals,
                 'SURAT KEPUTUSAN DIREKSI TENTANG PENETAPAN SISTEM MANAJEMEN PENGENDALIAN DOKUMEN DAN PENUNJUKAN TIM DOCUMENT CONTROL CENTER (DCC).',
                 "BOARD OF DIRECTORS' RESOLUTION CONCERNING THE ESTABLISHMENT OF THE DOCUMENT CONTROL MANAGEMENT SYSTEM AND APPOINTMENT OF THE DCC TEAM.",
                 is_bold=True)

    # KESATU
    add_dual_row(t_recitals, 'KESATU :', 'FIRST :', is_title=True, is_bold=True)
    add_dual_row(t_recitals,
                 'Memberlakukan secara resmi Standar Format dan Sistem Penomoran Dokumen 6 Segmen di seluruh lingkungan operasional dan kantor pusat PT Nirwana Bhumi Energy:\n[NomorUrut4Digit] - [Kategori] - [Divisi] - [Pengesah] - [BulanRomawi] - [Tahun]\nKetentuan penomoran ini bersifat unik mutlak: SATU KODE UNTUK SATU DOKUMEN (1 Kode Dokumen = 1 Dokumen Sah). Duplikasi nomor dilarang keras.',
                 'To officially enact the Standard Format and 6-Segment Document Numbering System across all operational areas and the head office of PT Nirwana Bhumi Energy:\n[4DigitSequentialNo] - [Category] - [Division] - [Authoriser] - [RomanMonth] - [Year]\nThis numbering provision is strictly unique: ONE CODE FOR ONE DOCUMENT (1 Document Code = 1 Valid Document). Duplication is strictly prohibited.')

    # Page Break to Page 4
    p_pb_h3 = doc.add_paragraph()
    p_pb_h3.add_run().add_break(WD_BREAK.PAGE)

    # --- PAGE 4: KEDUA, DCC TEAM TABLE, KETIGA, KEEMPAT, KELIMA, SIGNATURES ---
    t_body2 = doc.add_table(rows=0, cols=2)
    t_body2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_none(t_body2)

    add_dual_row(t_body2, 'KEDUA :', 'SECOND :', is_title=True, is_bold=True)
    add_dual_row(t_body2,
                 'Menunjuk dan menetapkan personel berikut sebagai Tim Document Control Center (DCC) PT Nirwana Bhumi Energy:',
                 'To appoint and designate the following personnel as the Document Control Centre (DCC) Team of PT Nirwana Bhumi Energy:')

    p_sp1 = doc.add_paragraph()
    p_sp1.paragraph_format.space_before = Pt(1)
    p_sp1.paragraph_format.space_after = Pt(1)

    # Tim DCC Table (Full Width 6.44 inches)
    table_team = doc.add_table(rows=4, cols=4)
    table_team.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = [
        'No.',
        'Nama Personel /\nPersonnel Name',
        'Jabatan Struktural /\nStructural Position',
        'Peran & Tanggung Jawab DCC /\nDCC Roles & Responsibilities'
    ]
    col_widths = [Inches(0.4), Inches(1.8), Inches(1.7), Inches(2.54)]

    for c_i, name in enumerate(headers):
        cell = table_team.cell(0, c_i)
        cell.width = col_widths[c_i]
        cell.paragraphs[0].text = name
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(cell, 'D9E1F2')
        set_cell_vcenter(cell)
        set_cell_margins(cell, top=35, bottom=35, left=50, right=50)
        for r in cell.paragraphs[0].runs:
            style_run(r, font_name='Arial', size_pt=7.5, bold=True, color_rgb=(0, 0, 0))

    team_data = [
        ('1.', 'Mazmur Gusti Agung Larosa', 'Direktur Keuangan & Teknik /\nDirector of Finance & Engineering', 
         'Lead Document Controller & System Administrator:\nMemvalidasi naskah, registrasi kode unik, eksekusi Fix Kode, dan sinkronisasi arsip Google Drive.\nValidates drafts, registers unique codes, executes Fix Code, and synchronises corporate Google Drive archives.'),
        ('2.', 'Dennis Raditya', 'Direktur Operasional & Teknologi /\nDirector of Operations & Technology', 
         'Technical Verifier & QA Approver:\nMemeriksa kesesuaian format teknis, layout, serta kepatuhan standar mutu IMS & CSMS.\nReviews technical format conformity, layout, and compliance with IMS & CSMS quality standards.'),
        ('3.', 'Abdul Qodir Al Jailani', 'Komisaris / IMS Lead /\nCommissioner / IMS Lead', 
         'Independent Compliance Overseer:\nMengawasi audit berkala dokumen terkontrol, penomoran masa lalu (backdating), dan integritas retensi arsip.\nSupervises periodic audits of controlled documents, past numbering (backdating), and archival retention integrity.')
    ]

    for r_i, data in enumerate(team_data, start=1):
        row = table_team.rows[r_i]
        set_row_cant_split(row)
        bg = 'FFFFFF' if r_i % 2 != 0 else 'F9FBFD'
        for c_i, val in enumerate(data):
            cell = row.cells[c_i]
            cell.width = col_widths[c_i]
            cell.paragraphs[0].text = val
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if c_i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_shading(cell, bg)
            set_cell_vcenter(cell)
            set_cell_margins(cell, top=30, bottom=30, left=50, right=50)
            for r in cell.paragraphs[0].runs:
                style_run(r, font_name='Arial', size_pt=7.0, bold=False, color_rgb=(0x3B, 0x38, 0x38))

    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_before = Pt(1)
    p_sp2.paragraph_format.space_after = Pt(1)

    # Remaining Clauses in Table 2-Col
    t_body3 = doc.add_table(rows=0, cols=2)
    t_body3.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_none(t_body3)

    add_dual_row(t_body3, 'KETIGA :', 'THIRD :', is_title=True, is_bold=True)
    add_dual_row(t_body3,
                 'Mengatur tata kelola penyimpanan dan registrasi dokumen digital melalui Tombol "Fix Kode", di mana dokumen otomatis terintegrasi ke Google Drive korporat (Folder ID: 1ny_1VhXfSaNrj_XdoOsC7V7BO-z0K8eF) pada subfolder divisi terkait.',
                 'To govern the digital storage and registration of corporate documents via the "Fix Code" button, whereby documents are automatically archived into corporate Google Drive (Folder ID: 1ny_1VhXfSaNrj_XdoOsC7V7BO-z0K8eF) within divisional subfolders.')

    add_dual_row(t_body3, 'KEEMPAT :', 'FOURTH :', is_title=True, is_bold=True)
    add_dual_row(t_body3,
                 'Menetapkan alur revisi dokumen terkontrol: penambahan sufiks (-rev1, dst.), pencatatan rincian alasan pada Revision Index, serta penarikan/pemusnahan salinan dokumen usang (obsolete).',
                 'To establish the revision workflow: suffix appending (-rev1, etc.), detailed recording in the Revision Index, and prompt withdrawal/destruction of obsolete copies.')

    add_dual_row(t_body3, 'KELIMA :', 'FIFTH :', is_title=True, is_bold=True)
    add_dual_row(t_body3,
                 'Surat Keputusan ini mulai berlaku efektif sejak tanggal ditetapkan, dan akan ditinjau kembali apabila terdapat penyesuaian regulasi industri migas/ketenagalistrikan.',
                 'This Resolution enters into force effectively from the enactment date, and will be re-evaluated should industry regulatory adjustments arise.')

    add_dual_row(t_body3,
                 'Ditetapkan di : Semarang / Tuban\nPada tanggal  : 17 Oktober 2026\nDIREKSI PT NIRWANA BHUMI ENERGY',
                 'Enacted in : Semarang / Tuban\nOn date    : 17 October 2026\nBOARD OF DIRECTORS OF PT NIRWANA BHUMI ENERGY',
                 is_bold=True)

    p_sp3 = doc.add_paragraph()
    p_sp3.paragraph_format.space_before = Pt(1)
    p_sp3.paragraph_format.space_after = Pt(1)

    # Signatures Table (3 columns)
    t_sig = doc.add_table(rows=3, cols=3)
    t_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_row_cant_split(t_sig.rows[0])
    set_row_cant_split(t_sig.rows[1])
    set_row_cant_split(t_sig.rows[2])

    sig_roles = [
        'Disetujui Oleh / Approved by:\nPresident Director / CEO',
        'Diverifikasi / Verified by:\nDirector of Operations & Tech (CTO)',
        'Disiapkan / Prepared by:\nDirector of Finance & Eng (CFO / DCC)'
    ]
    sig_names = [
        'BIMA SATRYO HARTANTO\nDirektur Utama / President Director',
        'DENNIS RADITYA\nDirektur Operasional / Director of Ops',
        'MAZMUR GUSTI AGUNG LAROSA\nDirektur Keuangan / Director of Finance'
    ]

    for c_i, role in enumerate(sig_roles):
        cell = t_sig.cell(0, c_i)
        cell.paragraphs[0].text = role
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_vcenter(cell)
        set_cell_margins(cell, top=30, bottom=30, left=45, right=45)
        for r in cell.paragraphs[0].runs:
            style_run(r, font_name='Arial', size_pt=7.5, bold=True, color_rgb=(0, 0, 0))

    for c_i in range(3):
        cell = t_sig.cell(1, c_i)
        cell.paragraphs[0].text = '\n[ Tanda Tangan & Cap ]\n' if c_i != 0 else '\n[ Meterai Rp 10.000 ]\n[ Tanda Tangan & Cap ]\n'
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_vcenter(cell)
        set_cell_margins(cell, top=35, bottom=35, left=45, right=45)
        for r in cell.paragraphs[0].runs:
            style_run(r, font_name='Arial', size_pt=7.0, italic=True, color_rgb=(0x88, 0x88, 0x88))

    for c_i, name in enumerate(sig_names):
        cell = t_sig.cell(2, c_i)
        cell.paragraphs[0].text = name
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_vcenter(cell)
        set_cell_margins(cell, top=30, bottom=30, left=45, right=45)
        for r in cell.paragraphs[0].runs:
            style_run(r, font_name='Arial', size_pt=7.5, bold=True, color_rgb=(0, 0, 0))

    doc.save(out_docx)
    print(f"Generated DOCX (Bilingual): {out_docx}")
    convert_docx_to_pdf(out_docx, out_pdf)
    return out_docx, out_pdf

if __name__ == '__main__':
    print("Building Bilingual Document...")
    bi_docx, bi_pdf = build_bilingual()

    print("Checking page count with PyMuPDF...")
    pdoc = pymupdf.open(bi_pdf)
    print(f"File: {os.path.basename(bi_pdf)} -> Total Pages: {len(pdoc)}")
